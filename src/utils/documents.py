"""Ingest documents users send over WhatsApp: download, extract text, chunk,
embed, store. Everything here is blocking; callers run it via asyncio.to_thread.
"""
import hashlib
import io
import logging
import re
import time
from dataclasses import dataclass
from urllib.parse import urlparse

import httpx

from src import database as db
from src.config import OPENAI_API_KEY, TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN

logger = logging.getLogger(__name__)

MAX_BYTES = 10 * 1024 * 1024
MEDIA_RETRY_PAUSES = [1, 2, 3, 5]  # seconds; ~11 s total before giving up on Twilio media
MAX_CHARS = 300_000
CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150
EMBED_MODEL = "text-embedding-3-small"
EMBED_BATCH = 96

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_TYPES = {"text/plain", "text/csv"}
SUPPORTED_TYPES = {PDF, DOCX} | TEXT_TYPES
_EXTENSIONS = {PDF: "pdf", DOCX: "docx", "text/plain": "txt", "text/csv": "csv"}


class DocumentError(Exception):
    """A problem the user should be told about in plain words."""


@dataclass
class IngestResult:
    status: str  # "saved" | "duplicate"
    document_id: int
    filename: str
    char_count: int
    chunk_count: int = 0
    searchable_by_meaning: bool = False
    supersedes_id: int | None = None


def _is_twilio_host(url: str) -> bool:
    host = (urlparse(url).hostname or "").lower()
    return host == "twilio.com" or host.endswith(".twilio.com")


def download_media(url: str) -> tuple[bytes, str | None]:
    """Fetch a Twilio media URL. Returns (bytes, filename-if-known).

    Twilio credentials are only ever sent to a *.twilio.com host, and
    redirects are followed by hand without them: Twilio answers media
    requests with a redirect to a storage host, and credentials must never
    ride along to whatever that (or a spoofed MediaUrl) points at.
    """
    auth = (TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN) if _is_twilio_host(url) else None
    # Some hosts refuse the default python-httpx agent outright (403).
    headers = {"User-Agent": "Mozilla/5.0 (compatible; Cue/1.0; +https://sellifyagent-production.up.railway.app)"}
    with httpx.Client(timeout=30.0, follow_redirects=False, headers=headers) as client:
        # Twilio fires the WhatsApp webhook before the media has always landed
        # in its store: an immediate GET can 404 and succeed a second later.
        # Retry with a short backoff rather than telling the user the photo
        # "failed to load".
        for attempt, pause in enumerate(MEDIA_RETRY_PAUSES + [0]):
            resp = client.get(url, auth=auth)
            hops = 0
            while resp.is_redirect and hops < 5:
                nxt = str(resp.next_request.url) if resp.next_request else resp.headers.get("location", "")
                hop_auth = auth if _is_twilio_host(nxt) else None
                resp = client.get(nxt, auth=hop_auth)
                hops += 1
            if resp.status_code in (404, 500, 502, 503) and auth and pause:
                logger.info("Media not ready yet (HTTP %s), retrying in %ss", resp.status_code, pause)
                time.sleep(pause)
                continue
            break
        resp.raise_for_status()
        data = resp.content
        if len(data) > MAX_BYTES:
            raise DocumentError(f"That file is too large ({len(data) // 1_000_000} MB). The limit is 10 MB.")
        match = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)"?', resp.headers.get("content-disposition", ""))
        return data, (match.group(1).strip() if match else None)


def extract_text(data: bytes, content_type: str) -> str:
    content_type = (content_type or "").split(";")[0].strip().lower()
    if content_type == PDF:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError("That PDF is password-protected, so I can't read it.")
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except DocumentError:
            raise
        except Exception as e:
            raise DocumentError(f"I couldn't open that PDF ({type(e).__name__}).")
        if len(text.strip()) < 20:
            # Scanned PDFs are images of pages with no text layer; reading
            # them needs OCR/vision, which isn't wired up yet.
            raise DocumentError("That PDF looks like a scan (images, no text), which I can't read yet.")
        return text
    if content_type == DOCX:
        from docx import Document

        try:
            doc = Document(io.BytesIO(data))
        except Exception as e:
            raise DocumentError(f"I couldn't open that Word file ({type(e).__name__}).")
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    if content_type in TEXT_TYPES:
        return data.decode("utf-8", errors="replace")
    raise DocumentError(
        "I can't read that file type yet. I can read PDF, Word (.docx) and plain-text files."
    )


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split on paragraph boundaries into ~size-char chunks, with a small
    overlap so a fact spanning a boundary is still findable in one chunk."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(para[:size])
            para = para[size - overlap:]
        if len(current) + len(para) + 2 <= size:
            current = f"{current}\n\n{para}" if current else para
        else:
            if current:
                chunks.append(current)
                tail = current[-overlap:]
                current = f"{tail}\n\n{para}"
            else:
                current = para
    if current:
        chunks.append(current)
    return chunks


def embed(texts: list[str]) -> list[list[float]] | None:
    """OpenAI embeddings, or None when no key is configured (keyword search
    still works then). Raises on API errors so a half-embedded document is
    never stored."""
    if not OPENAI_API_KEY:
        return None
    vectors: list[list[float]] = []
    with httpx.Client(timeout=60.0) as client:
        for i in range(0, len(texts), EMBED_BATCH):
            resp = client.post(
                "https://api.openai.com/v1/embeddings",
                headers={"Authorization": f"Bearer {OPENAI_API_KEY}"},
                json={"model": EMBED_MODEL, "input": texts[i:i + EMBED_BATCH]},
            )
            resp.raise_for_status()
            data = sorted(resp.json()["data"], key=lambda d: d["index"])
            vectors.extend(d["embedding"] for d in data)
    return vectors


def embed_query(query: str) -> list[float] | None:
    try:
        vectors = embed([query])
        return vectors[0] if vectors else None
    except Exception:
        logger.warning("Query embedding failed; falling back to keyword search", exc_info=True)
        return None


def default_filename(content_type: str, message_sid: str | None) -> str:
    ext = _EXTENSIONS.get((content_type or "").split(";")[0].strip().lower(), "bin")
    suffix = (message_sid or "")[-6:] or "upload"
    return f"document-{suffix}.{ext}"


def ingest(user_id: str, data: bytes, content_type: str, filename: str) -> IngestResult:
    content_hash = hashlib.sha256(data).hexdigest()
    existing = db.find_document_by_hash(user_id, content_hash)
    if existing:
        return IngestResult("duplicate", existing["id"], existing["filename"], 0)

    text = extract_text(data, content_type)
    if len(text) > MAX_CHARS:
        raise DocumentError(
            f"That document is very long ({len(text):,} characters). The limit is {MAX_CHARS:,}."
        )
    chunks = chunk_text(text)
    if not chunks:
        raise DocumentError("That document has no readable text.")

    embeddings = embed(chunks)
    saved = db.save_document(
        user_id, filename, content_type, len(data), content_hash, chunks, embeddings
    )
    logger.info("Ingested document %s for user (%d chunks, embedded=%s)", saved["id"], len(chunks), bool(embeddings))
    return IngestResult(
        "saved",
        saved["id"],
        filename,
        sum(len(c) for c in chunks),
        chunk_count=len(chunks),
        searchable_by_meaning=bool(embeddings),
        supersedes_id=saved["supersedes_id"],
    )
